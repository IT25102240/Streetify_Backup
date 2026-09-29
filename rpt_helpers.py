"""
rpt_helpers.py  — Shared helper functions for Streetify report generator
"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── Colors ────────────────────────────────────────────────────────────────────
C_DARK  = RGBColor(0x0D, 0x2B, 0x55)
C_BLUE  = RGBColor(0x1A, 0x56, 0xDB)
C_ORNG  = RGBColor(0xF9, 0x73, 0x16)
C_GREY  = RGBColor(0x37, 0x41, 0x51)
C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
C_CODE  = RGBColor(0xAB, 0xB2, 0xBF)
C_GREEN = RGBColor(0x05, 0x96, 0x69)

# ── XML helpers ───────────────────────────────────────────────────────────────
def shd_cell(cell, hx):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hx)
    tcPr.append(shd)

def shd_para(p, hx):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hx)
    pPr.append(shd)

def add_hr(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(2)
    pPr = p._p.get_or_add_pPr(); pBdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "6")
    b.set(qn("w:space"), "1"); b.set(qn("w:color"), "1A56DB")
    pBdr.append(b); pPr.append(pBdr)

# ── Content helpers ───────────────────────────────────────────────────────────
def H(doc, txt, lv=1):
    h = doc.add_heading(txt, level=lv)
    for r in h.runs:
        r.font.color.rgb = C_DARK if lv == 1 else C_BLUE
        r.font.bold = True

def P(doc, txt="", bold=False, italic=False, sz=11, color=None, sa=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(sa)
    if txt:
        r = p.add_run(txt); r.bold = bold; r.italic = italic; r.font.size = Pt(sz)
        if color: r.font.color.rgb = color
    return p

def CODE(doc, code, title=None):
    if title:
        tp = doc.add_paragraph()
        tp.paragraph_format.space_before = Pt(8); tp.paragraph_format.space_after = Pt(0)
        tr = tp.add_run("  " + title)
        tr.font.name = "Consolas"; tr.font.size = Pt(9); tr.bold = True; tr.font.color.rgb = C_ORNG
        shd_para(tp, "1E1E2E")
    for line in code.split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(0.3)
        r = p.add_run(line if line else " ")
        r.font.name = "Consolas"; r.font.size = Pt(9.5); r.font.color.rgb = C_CODE
        shd_para(p, "1E1E2E")
    sp = doc.add_paragraph(" "); sp.paragraph_format.space_after = Pt(10)

def BOX(doc, txt, t="note"):
    cols = {"note": ("2B6CB0", "NOTE: "), "tip": ("276749", "VIVA TIP: "), "warn": ("9C4221", "IMPORTANT: ")}
    hx, pre = cols.get(t, cols["note"])
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5); p.paragraph_format.space_after = Pt(8)
    r1 = p.add_run(pre); r1.bold = True; r1.font.size = Pt(10.5)
    r1.font.color.rgb = RGBColor(int(hx[0:2], 16), int(hx[2:4], 16), int(hx[4:6], 16))
    r2 = p.add_run(txt); r2.font.size = Pt(10.5); r2.font.color.rgb = C_GREY
    shd_para(p, "EBF4FF" if t == "note" else ("EBFFF4" if t == "tip" else "FFF3EB"))

def TBL(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h; shd_cell(hdr[i], "0D2B55")
        for pa in hdr[i].paragraphs:
            for r in pa.runs: r.font.bold = True; r.font.size = Pt(10); r.font.color.rgb = C_WHITE
            pa.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for ri, row in enumerate(rows):
        cells = t.add_row().cells; bg = "EBF2FF" if ri % 2 == 0 else "FFFFFF"
        for ci, val in enumerate(row):
            cells[ci].text = val; shd_cell(cells[ci], bg)
            for pa in cells[ci].paragraphs:
                for r in pa.runs: r.font.size = Pt(9.5); r.font.name = "Calibri"
    doc.add_paragraph()

def QA(doc, q, a):
    p1 = doc.add_paragraph(); p1.paragraph_format.left_indent = Cm(0.5)
    p1.paragraph_format.space_before = Pt(6); p1.paragraph_format.space_after = Pt(2)
    r1 = p1.add_run("Q: " + q); r1.bold = True; r1.font.color.rgb = C_DARK; r1.font.size = Pt(10.5)
    p2 = doc.add_paragraph(); p2.paragraph_format.left_indent = Cm(1.0); p2.paragraph_format.space_after = Pt(8)
    r2 = p2.add_run("A: " + a); r2.font.color.rgb = C_GREY; r2.font.size = Pt(10.5)

def SCRN(doc, label):
    p = doc.add_paragraph()
    r = p.add_run(f"[ SCREENSHOT: {label} ]")
    r.bold = True; r.font.color.rgb = C_ORNG; r.font.size = Pt(10)
    p.paragraph_format.space_after = Pt(12)
